"""Build de-duplicated Islamabad and Rawalpindi service-area lists from the owner's master spreadsheets.

Usage:
    python build_area_lists.py <rawalpindi.xlsx> <islamabad.xlsx>

Writes next to this script:
    islamabad_areas.csv, rawalpindi_areas.csv   one row per area, each area in one city only
    overlap_decisions.csv                       every area found in both files and where it was kept
Edit OVERRIDES below and re-run to move an area to the other city.
"""
import csv
import os
import re
import sys

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROMAN = {"i": "1", "ii": "2", "iii": "3", "iv": "4", "v": "5", "vi": "6", "vii": "7", "viii": "8", "ix": "9"}

# Name in both files -> city it stays in. "both" = same name, but two different places (one in each city).
OVERRIDES = {
    # Islamabad-marketed schemes that sit in Rawalpindi district
    "Al-Haram City": "rawalpindi", "Bahria Intellectual Village": "rawalpindi", "Blue World City": "rawalpindi",
    "Capital Smart City": "rawalpindi", "Faisal Hills": "rawalpindi", "Kingdom Valley": "rawalpindi",
    "Mumtaz City": "rawalpindi", "Rudn Enclave": "rawalpindi", "Silver City": "rawalpindi",
    "Taj Residencia": "rawalpindi", "Gandhara City": "rawalpindi", "Media Town": "rawalpindi",
    "Airport Enclave": "rawalpindi", "Airport Green Garden": "rawalpindi",
    # Rawalpindi city / cantonment localities and roads
    "Dhok Kala Khan": "rawalpindi", "Shakrial": "rawalpindi", "Rawat": "rawalpindi",
    "Murree Road": "rawalpindi", "Faizabad Interchange": "rawalpindi",
    # CDA-listed or ICT schemes
    "Army Welfare Trust Housing Scheme": "islamabad", "Bahria Garden City": "islamabad",
    "T&T Employees Cooperative Housing Society": "islamabad", "Korang Town": "islamabad",
    "Pakistan Town": "islamabad", "Faisal Margalla City": "islamabad", "Top City-1": "islamabad",
    "Foreign Office Employees Cooperative Housing Society": "islamabad", "Lehtrar Road": "islamabad",
    "Shalimar Town": "islamabad",
    # Common names used for two unrelated places (e.g. Malpur near Margalla vs Malpur in Taxila)
    "Awan Town": "both", "Azam Colony": "both", "Madina Colony": "both", "Doctors Colony": "both",
    "Officers Colony": "both", "Gulshan-e-Abad": "both", "Dhok Syedan": "both", "Mohra Nagiyal": "both",
    "Malpur": "both", "Kuri": "both", "Shahpur": "both",
}


def key(name):
    s = name.lower().replace("&", " and ")
    s = re.sub(r"\b(islamabad|rawalpindi)\b", " ", s)
    s = re.sub(r"\bmulti gardens\b", " ", s)  # "Multi Gardens B-17 Block A" == "B-17 Block A"
    s = re.sub(r"\b(phase|ph)[\s-]*([ivx]+)\b", lambda m: "phase " + ROMAN.get(m.group(2), m.group(2)), s)
    return re.sub(r"[^a-z0-9]", "", s)


def family_city(k):
    """Whole society families that appear in both files under different spellings."""
    if k.startswith("bahriatownphase"):
        return "rawalpindi"
    # Match DHA society names only; a bare "dha" prefix also caught villages such as Dhamial and Dhari Syedan.
    if re.match(r"dha(phase|valley|sector|\d)", k) or k.startswith("b17"):
        return "islamabad"
    return None


def read(path, sheet, cols):
    rows = list(load_workbook(path, read_only=True)[sheet].iter_rows(values_only=True))[1:]
    out, seen = [], set()
    for r in rows:
        if not r[0] or key(str(r[0])) in seen:
            continue
        seen.add(key(str(r[0])))
        out.append({c: (r[i] or "") for i, c in enumerate(cols)})
    return out


def main(rwp_path, isb_path):
    cols = ["Area", "Group", "Type", "Notes", "Source"]
    rwp = read(rwp_path, "All Areas", ["Area", "Group", "Type", "Parent", "Notes", "Source"])
    isb = read(isb_path, "All locations", ["Area", "Type", "Group", "Jurisdiction", "Notes", "Source"])
    rwp_by_key = {key(r["Area"]): r for r in rwp}
    isb_by_key = {key(r["Area"]): r for r in isb}

    decisions = []
    for k in sorted(set(rwp_by_key) & set(isb_by_key)):
        name = rwp_by_key[k]["Area"]
        city = OVERRIDES.get(name) or family_city(k)
        if not city:
            raise SystemExit(f"No decision for '{name}' (in both files): add it to OVERRIDES")
        decisions.append({"Area": name, "Islamabad file name": isb_by_key[k]["Area"],
                          "Kept in": {"both": "Both (different places)"}.get(city, city.title())})
        if city == "islamabad":
            rwp.remove(rwp_by_key[k])
        elif city == "rawalpindi":
            isb.remove(isb_by_key[k])

    # Sub-areas of those families that only one file lists ("DHA Phase 1 Sector A") follow their parent
    for row in [r for r in rwp if family_city(key(r["Area"])) == "islamabad"]:
        rwp.remove(row)
        isb.append(dict(row, Notes=f"Moved from the Rawalpindi list. {row['Notes']}".strip()))
    for row in [r for r in isb if family_city(key(r["Area"])) == "rawalpindi"]:
        isb.remove(row)
        rwp.append(dict(row, Notes=f"Moved from the Islamabad list. {row['Notes']}".strip()))

    for name, rows in (("islamabad", isb), ("rawalpindi", rwp)):
        rows.sort(key=lambda r: (str(r["Group"]).lower(), str(r["Area"]).lower()))
        with open(os.path.join(HERE, f"{name}_areas.csv"), "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
            w.writeheader()
            w.writerows(rows)
    with open(os.path.join(HERE, "overlap_decisions.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["Area", "Islamabad file name", "Kept in"])
        w.writeheader()
        w.writerows(decisions)

    counts = {c: sum(1 for d in decisions if d["Kept in"].startswith(c)) for c in ("Islamabad", "Rawalpindi", "Both")}
    print(f"Islamabad {len(isb)} areas, Rawalpindi {len(rwp)} areas; overlaps {len(decisions)}: {counts}")


if __name__ == "__main__":
    main(*sys.argv[1:3])
