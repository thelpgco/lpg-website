import openpyxl, re, json, csv, collections as C
SRC=r'D:\SparePartsHub-App\lpg-website\keywords and research\LPG_dealers_all_regions.xlsx'
ws=openpyxl.load_workbook(SRC,data_only=True).active
hdr=[c.value for c in ws[1]]
rows=[dict(zip(hdr,[c.value for c in r])) for r in ws.iter_rows(min_row=2)]
PROV={
'Sindh':['karachi','hyderabad','sukkur','larkana','nawabshah','mirpur khas','khairpur','jacobabad','shikarpur','badin','thatta','dadu','sanghar','tando','thul','shah faisal','jamshoro','kotri','ghotki','umerkot','moro','hala','matli','rohri','daharki','kandhkot','mehar','sehwan','naushahro','gambat','shahdadpur','kamber','lakhi','pano akil','ratodero','sakrand'],
'Punjab':['lahore','rawalpindi','faisalabad','multan','gujranwala','sialkot','sargodha','bahawalpur','sahiwal','sheikhupura','gujrat','jhelum','chakwal','attock','kasur','okara','rahim yar','dera ghazi','mianwali','jhang','chiniot','vehari','khanewal','muzaffargarh','gojra','burewala','bhakkar','talagang','shujaab','pasr','taxila','wah','gujar khan','kharian','hafizabad','mandi bahauddin','daska','pakpattan','bahawalnagar','kamoke','kāmoke','murid','lodhran','layyah','narowal','wazirabad','jauharabad','khushab','nankana','hasilpur','chishtian','haroonabad','jaranwala','kot addu','mian channu','lala musa','sambrial','bhalwal','pattoki','toba','samundri','thakkarke','qalandar','fateh jang','hasan abdal','kahuta','murree','chunian','dinga','liaqatpur','arifwala','arif wala','sadiqabad','khanpur','ahmedpur','ferozwala','muridke','kot abdul','jampur','rajanpur','mailsi','jatoi','kabirwala','shakargarh','pindi','bhera','kallar','chak','shahkot','zafarwal','sangla','lalian','dunyapur','jahanian'],
'KP':['peshawar','mardan','abbottabad','mansehra','haripur','mingora','swat','kohat','nowshera','charsadda','bannu','dera ismail','dir','swabi','chitral','batkhela','timergara','shinkiari','jamrud','lakki','katair','talaab','thal','tehsil munda','osakai','raporai','topi','thakot','thana','karak','hangu','kabal','takht','barikot','shabqadar','jehangira','tank','buner','malakand','kurram','parachinar','tordher','sawabi','daggar','oghi','battagram','balakot'],
'Balochistan':['quetta','gwadar','hub','khuzdar','chaman','turbat','sibi','zhob','loralai','dera murad','harnai','nasirabad','pishin','panjgur','kharan','usta','jaffarabad','mastung','kalat','ziarat','dera allah','sui','bela','lasbela','uthal'],
'AJK':['muzaffarabad','mirpur','rawalakot','kotli','bhimber','new mirpur','bagh','neelum','poonch','sudhnoti','haveli','athmuqam'],
'GB':['gilgit','skardu','chilas','astore','hunza','ali abad','ghizer','diamer','ghanche','khaplu','nagar','aliabad'],
'ICT':['islamabad'],
}
def prov(city,addr):
    t=((city or '')+' | '+(addr or '')).lower()
    c=(city or '').lower()
    for p,keys in PROV.items():
        if any(k in c for k in keys): return p
    for p,keys in PROV.items():
        if any(k in t for k in keys): return p
    return 'UNKNOWN'
def norm_phone(p):
    d=re.sub(r'\D','',p or '')
    if d.startswith('92'): d=d[2:]
    if d.startswith('0'): d=d[1:]
    return d
GOOD={'Gas cylinders supplier','Gas company','Gas shop','Gas station','Wholesaler','Distribution service','Fuel supplier','Energy supplier','Alternative fuel station','Oil & natural gas company','Petroleum products company','Warehouse','Corporate office','Manufacturer','Transportation service','Store'}
def classify(r):
    n=(r['name'] or '').lower(); st=(r['subtypes'] or '')
    types=[x.strip() for x in st.split(',') if x.strip()]
    conv=bool(re.search(r'\b(cng|conversion|auto ?gas|kit|car|motors?|auto|workshop|filling station|petrol|pump|oil)\b',n)) or 'LPG conversion' in types
    cyl=bool(re.search(r'\b(lpg|gas|cylinder)\b',n)) and not re.search(r'\bcng\b|conversion',n)
    if types==['LPG conversion'] or (conv and not any(t in ('Gas cylinders supplier','Gas shop') for t in types)):
        if re.search(r'\b(lpg)\b',n) and re.search(r'(distribut|cylinder|agency|supplier|gas shop|filling plant|dealer|wholesale)',n):
            return 'REVIEW'
        return 'AUTO_CONVERSION'
    if any(t in ('Gas cylinders supplier','Gas shop','Wholesaler','Distribution service','Gas company') for t in types) or cyl: return 'LPG_DEALER'
    if any(t in ('Gas station','Fuel supplier','Alternative fuel station','Petroleum products company','Oil & natural gas company','Energy supplier') for t in types): return 'REVIEW'
    return 'OTHER'
out=[];seen={}
dropped=C.Counter()
for r in rows:
    r['phone_n']=norm_phone(r['phone'])
    r['province']=prov(r['city'],r['address'])
    r['class']=classify(r)
    k=(r['phone_n'], re.sub(r'[^a-z0-9]','',(r['name'] or '').lower())[:18])
    if r['phone_n'] in seen:
        dropped['dup_phone']+=1; continue
    seen[r['phone_n']]=1
    if not re.fullmatch(r'3\d{9}|[1-9]\d{8,10}',r['phone_n']): r['class']+='|BAD_PHONE'
    if r['website'] and re.search(r'facebook|instagram|youtube|business\.site',r['website']): r['website_type']='social_or_free'
    elif r['website']: r['website_type']='own_site'
    else: r['website_type']=''
    out.append(r)
print('in',len(rows),'after dedupe',len(out),dropped)
print(C.Counter(r['class'] for r in out))
print(C.Counter(r['province'] for r in out))
print('\nLPG_DEALER by province',C.Counter(r['province'] for r in out if r['class']=='LPG_DEALER'))
print('\nUNKNOWN',[(r['city'],(r['address'] or '')[-45:]) for r in out if r['province']=='UNKNOWN'][:30])
print('\nLPG_DEALER per city (top)',C.Counter(r['city'] for r in out if r['class']=='LPG_DEALER').most_common(25))
print('own sites among LPG_DEALER',sum(1 for r in out if r['class']=='LPG_DEALER' and r['website_type']=='own_site'))
print('REVIEW sample',[r['name'] for r in out if r['class']=='REVIEW'][:15])
print('OTHER sample',[r['name'] for r in out if r['class']=='OTHER'][:12])
cols=['name','name_for_emails','subtypes','phone','phone_n','website','website_type','address','city','province','class']
with open('../data/dealers_xlsx_classified.csv','w',newline='',encoding='utf-8-sig') as f:
    w=csv.DictWriter(f,fieldnames=cols,extrasaction='ignore');w.writeheader();w.writerows(out)

