import openpyxl, re, collections as C
wb=openpyxl.load_workbook('LPG_dealers_all_regions.xlsx',data_only=True)
ws=wb.active
rows=[dict(zip([c.value for c in ws[1]],[c.value for c in r])) for r in ws.iter_rows(min_row=2)]
print('rows',len(rows))
print('blank-name',sum(1 for r in rows if not r['name']))
print('phone missing',sum(1 for r in rows if not r['phone']))
print('website present',sum(1 for r in rows if r['website']))
print('address missing',sum(1 for r in rows if not r['address']))
print('city missing',sum(1 for r in rows if not r['city']))
print('country',C.Counter(r['country'] for r in rows))
print('\nSUBTYPES');[print(' ',k,v) for k,v in C.Counter(r['subtypes'] for r in rows).most_common(25)]
print('\nCITIES',len(set(r['city'] for r in rows)));[print(' ',k,v) for k,v in C.Counter(r['city'] for r in rows).most_common(60)]
def norm(p):
    d=re.sub(r'\D','',p or '')
    if d.startswith('92'): d=d[2:]
    if d.startswith('0'): d=d[1:]
    return d
pc=C.Counter(norm(r['phone']) for r in rows if r['phone'])
print('\ndup phones',sum(1 for k,v in pc.items() if v>1), 'rows in dup',sum(v for v in pc.values() if v>1))
nc=C.Counter((r['name'] or '').strip().lower() for r in rows)
print('dup names',sum(1 for k,v in nc.items() if v>1))
print('phone formats',C.Counter(re.sub(r'\d','9',r['phone'] or '') for r in rows).most_common(8))
print('\nwebsites sample',[r['website'] for r in rows if r['website']][:15])
print('\nweird city (not in address)',[(r['city'],r['address'][:60]) for r in rows if r['address'] and r['city'] and r['city'].lower() not in r['address'].lower()][:15])
